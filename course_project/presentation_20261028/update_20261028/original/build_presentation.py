"""Build the editable, evidence-grounded ten-minute class presentation."""
from pathlib import Path
import json
import math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[1]
NEW=json.loads((ROOT/'reports/rewt_reasoning_20261026/summary.json').read_text())
OLD=json.loads((ROOT/'reports/rewt_architecture_20261022/summary.json').read_text())
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
prs.core_properties.title='REWT: Learning an Energy Landscape for State Refinement'
prs.core_properties.subject='CSCI 635, ten-minute technical class presentation, October 28, 2026'
prs.core_properties.author='Kenju Tomita'
prs.core_properties.keywords='REWT, energy-based models, recurrent memory, learned reasoning, controlled experiments'
BG='F6F4EE';INK='172C45';TEAL='087E80';ORANGE='D86A32';GRAY='647282';LIGHT='E6EBEB';WHITE='FFFFFF';RED='A54D49'
notes=[]

def rgb(c):return RGBColor.from_string(c)
def rect(s,x,y,w,h,color,rounded=False,line=None):
    q=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    q.fill.solid();q.fill.fore_color.rgb=rgb(color)
    q.line.color.rgb=rgb(line or color)
    if rounded:q.adjustments[0]=.08
    return q
def txt(s,text,x,y,w,h,size=22,color=INK,bold=False,align=PP_ALIGN.LEFT,font='Arial'):
    shape=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=shape.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.01);tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.alignment=align
        p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color)
        p.space_after=Pt(5)
    return shape
def arrow(s,x,y,w=.5,h=.25,color=TEAL):
    return rect_arrow(s,x,y,w,h,color)
def rect_arrow(s,x,y,w,h,color):
    q=s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Inches(x),Inches(y),Inches(w),Inches(h))
    q.fill.solid();q.fill.fore_color.rgb=rgb(color);q.line.fill.background();return q
def line(s,x1,y1,x2,y2,color=GRAY,width=1.5):
    q=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    q.line.color.rgb=rgb(color);q.line.width=Pt(width);return q
def circle(s,label,x,y,diam=1,color=TEAL,fg=WHITE):
    q=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(diam),Inches(diam))
    q.fill.solid();q.fill.fore_color.rgb=rgb(color);q.line.fill.background()
    txt(s,label,x,y+.20*diam,diam,diam*.70,min(30,30*diam),fg,True,PP_ALIGN.CENTER)
    return q
def source(s,text):txt(s,text,.55,7.08,11.55,.24,10,GRAY)
def slide(kicker,title,subtitle='',source_text=''):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(BG)
    n=len(prs.slides)
    rect(s,.55,.44,.38,.07,TEAL)
    txt(s,kicker.upper(),1.04,.34,11,.3,11,TEAL,True)
    txt(s,title,.55,.92,12.2,.92,32,INK,True)
    if subtitle:txt(s,subtitle,.57,1.95,12.1,.55,18,GRAY)
    txt(s,f'{n:02d}',12.13,7.03,.6,.3,11,GRAY,False,PP_ALIGN.RIGHT)
    if source_text:source(s,source_text)
    return s
def note(s,seconds,body,refs=''):
    n=len(prs.slides)
    text=f'SLIDE {n} · {seconds} seconds\n\n{body}\n\nSources / evidence:\n{refs}'
    s.notes_slide.notes_text_frame.text=text
    notes.append(dict(slide=n,seconds=seconds,text=text))
def card(s,kicker,headline,body,x,y,w=3.9,h=2.8,color=TEAL):
    rect(s,x,y,w,h,WHITE,True)
    txt(s,kicker,x+.23,y+.22,w-.46,.4,13,color,True)
    txt(s,headline,x+.23,y+.8,w-.46,.8,26,INK,True)
    txt(s,body,x+.23,y+1.72,w-.46,h-1.85,18,GRAY)

# 1
s=slide('CSCI 635 · Research progress · October 28, 2026','')
txt(s,'REWT',.62,1.25,8,1.1,68,INK,True)
txt(s,'Learning an energy landscape\nfor state refinement',.65,2.55,10.8,1.5,32,INK)
txt(s,'Can revising an internal state improve an answer?',.67,4.55,11.8,.65,25,TEAL,True)
txt(s,'Kenju Tomita  ·  Rochester Institute of Technology',.68,6.25,11,.45,17,GRAY)
circle(s,'?',10.95,1.25,1.2,TEAL)
source(s,'Recurrent Energy Wave Transformer · Project evidence through October 26, 2026')
note(s,25,'My project asks whether a model can improve an answer by revising its internal state before emitting it. REWT stands for Recurrent Energy Wave Transformer. I will explain energy and waves intuitively, then separate what the experiments establish from what is still unproven. The strongest results so far concern learned memory and a new repeated-read module.','Project implementation: rewt/model.py, rewt/architecture.py, rewt/reasoning.py')

# 2
s=slide('Motivation','The goal: compose facts into an answer','A small rule task makes each reasoning step checkable.')
rect(s,.7,2.7,11.9,.78,INK,True)
txt(s,'B>C;  A>B;  C>D;  X>Y;  from A go 3=?',.98,2.91,11.3,.42,24,WHITE,False,font='DejaVu Sans Mono')
for j,label in enumerate('ABCD'):
    x=1.45+j*3.0;circle(s,label,x,4.15,1.0,TEAL if j<3 else ORANGE)
    if j<3:
        arrow(s,x+1.25,4.5,1.1,.27)
        txt(s,f'read {j+1}',x+1.15,5.03,1.3,.35,15,GRAY,False,PP_ALIGN.CENTER)
txt(s,'Correct trace: B, C, D      Final answer: D',1.02,6.13,11.2,.5,24,INK,True)
source(s,'Illustrative task, not a benchmark sample. Actual tests use shuffled facts and disjoint semantic identities.')
note(s,40,'Suppose the input says A leads to B, B to C and C to D, with distractors and shuffled order. To move three times from A, the model needs to use B as its next query and then C. We can score both the intermediate sequence and the final answer. This is a deliberately narrow test of composition. It is useful because a correct-looking final answer cannot hide a wrong intermediate trace.','Task generators: rewt/tasks.py, rewt/curriculum.py; new study: scripts/study_rewt_reasoning.py')

# 3
s=slide('Intuition','Energy is a learned compatibility score','Lower energy means the model prefers that state. It does not certify truth.')
s.shapes.add_picture(str(HERE/'energy_landscape.png'),Inches(.72),Inches(2.72),width=Inches(8.6),height=Inches(3.34))
txt(s,'energy',.76,2.45,1.0,.35,17,GRAY)
txt(s,'Anchor',2.2,2.45,1.5,.32,15,GRAY)
txt(s,'Anchor + waves',3.8,2.45,2.5,.32,15,TEAL,True)
txt(s,'candidate state',3.75,6.02,3,.38,17,GRAY)
card(s,'THE LOOP','Guess → revise','Follow the slope.\nAccept a lower score.\nRepeat with fixed weights.',9.55,2.65,3.1,3.6)
txt(s,'A lower local valley can still correspond to a wrong answer.',1.03,6.57,11.7,.4,19,ORANGE,True)
source(s,'Illustrative one-dimensional landscape; the implemented state is high-dimensional. Background: Hoover et al., 2023.')
note(s,55,'Energy here is just a scalar compatibility score, not physical energy. Imagine the current state as a ball on this landscape. A gradient tells it which direction lowers the score. Backtracking rejects a proposed update that raises the current energy. The landscape is learned, so a low valley can still represent a wrong answer. Also, gradient descent may reach a local valley rather than the global minimum. This drawing is a one-dimensional illustration; the actual joint state has many coordinates.','Energy Transformer: https://arxiv.org/abs/2302.07253\nImplementation: rewt/model.py, energy(), gradient(), refine(). Plot is illustrative, generated locally.')

# 4
s=slide('Architecture','One energy couples several state blocks','Observed text proposes an initial state; refinement updates the joint workspace.')
for j,(a,b) in enumerate([('Recurrent','past context'),('Attention','retrieved facts'),('Feedforward','local proposal'),('Output','next-byte logits')]):
    x=.67+j*3.18
    rect(s,x,2.8,2.85,1.35,WHITE,True)
    txt(s,a,x+.12,3.01,2.61,.4,22,INK,True,PP_ALIGN.CENTER)
    txt(s,b,x+.12,3.57,2.61,.3,16,GRAY,False,PP_ALIGN.CENTER)
    line(s,x+1.425,4.16,x+1.425,4.75,TEAL,2)
line(s,2.095,4.75,11.635,4.75,TEAL,2)
rect(s,2.1,5.0,9.55,.72,TEAL,True)
txt(s,'Energy = stabilizing quadratic + learned cosine waves',2.23,5.19,9.29,.38,22,WHITE,True,PP_ALIGN.CENTER)
txt(s,'Quadratic: stay near a context-informed proposal.   Waves: learn coupled preferences.',.85,6.2,11.65,.7,20,INK)
source(s,'Implemented language-model core. “Wave” means Fourier components in state space, not sound or a quantum mechanism.')
note(s,55,'The language-model core has recurrent, attention, feedforward and output blocks. Proposal networks build the initial guess from observed context. One scalar energy then scores the joint workspace. Its positive quadratic term anchors the state, while learned cosine components can couple different blocks. A wave is simply a Fourier component over state coordinates. This is different from applying a Fourier transform across the token sequence. The implementation is a small research model; calling it energy-based does not imply that every useful behavior is caused by the wave term.','Core implementation: rewt/model.py\nDerivations: reports/rewt_architecture_20261022/DESIGN.md')

# 5
s=slide('Learning','Training shapes the landscape. Inference moves the state.')
card(s,'TRAINING','Change parameters θ','Use labeled examples.\nDifferentiate through refinement.\nUpdate weights with AdamW.',.75,2.55,5.7,3.45,TEAL)
card(s,'INFERENCE','Change candidate state s','Read the supplied context.\nTake accepted state updates.\nKeep learned weights fixed.',6.85,2.55,5.7,3.45,ORANGE)
txt(s,'s ← s − η ∇Eθ(context, s)',3.0,6.25,7.4,.6,29,INK,True,PP_ALIGN.CENTER)
source(s,'State optimization and parameter optimization are separate. Fixed-weight and analytic-gradient checks are automated.')
note(s,50,'There are two different optimization problems. During training, the loss compares predictions with labels and updates the parameters that define proposals and the energy landscape. During inference, those weights stay fixed. Only the candidate state changes. The state update uses the derivative of energy with respect to the state, not an optimizer step on the model weights. In the experiments I verify that inference leaves the checkpoint tensors unchanged. A separate numerical test checks the hand-derived state gradient against automatic differentiation.','Training and inference: rewt/model.py, rewt/reasoning.py\nVerification: tests_rewt/test_reasoning.py; reports/rewt_reasoning_20261026/verification.json')

# 6
s=slide('Evaluation','Three questions need three different measurements')
card(s,'01 · NUMERICS','Does energy fall?','Check accepted updates and\nanalytic gradients.',.7,2.6,3.8,3.25)
card(s,'02 · QUALITY','Is the answer right?','Score the complete trace\nand final answer.',4.77,2.6,3.8,3.25)
card(s,'03 · COST','What does it cost?','Measure training and inference;\nkeep compute budgets visible.',8.84,2.6,3.8,3.25)
rect(s,.85,6.22,11.65,.58,LIGHT,True)
txt(s,'Disjoint problem identities  ·  Validation selection  ·  Three seeds  ·  Frozen test models',1.0,6.37,11.35,.3,17,INK,True,PP_ALIGN.CENTER)
source(s,'Protocols and complete outputs are saved in reports/rewt_architecture_20261022 and reports/rewt_reasoning_20261026.')
note(s,40,'A numerical guarantee is not an accuracy result. I separately check descent, complete answers and cost. The studies keep semantic problem identities separate across training, validation and test, so changing the wording cannot turn a training problem into a test example. Checkpoints are selected on validation. The fresh tests run after the declared fits are complete. Three initialization seeds reveal some training variation, although three seeds and synthetic tasks are still a limited basis for general claims.','Study plans and source/data checksums: reports/rewt_architecture_20261022/plan.json and reports/rewt_reasoning_20261026/plan.json')

# 7
s=slide('Earlier result · Persistent memory','Learned overwrite memory survives longer gaps','Return the latest value after a key has been written twice.')
txt(s,'K=2   …   K=9   …   query K → 9',1.0,2.62,11.2,.6,27,INK,True,PP_ALIGN.CENTER)
for y,label,key,col in [(3.75,'Local + binding/copy','local_binding',GRAY),(4.95,'Delta + binding/copy/write','delta_binding',TEAL)]:
    value=OLD['memory_ablation'][key]['test_mean']
    txt(s,label,.85,y,4.5,.42,21,INK)
    rect(s,5.4,y,5.5,.44,LIGHT,True);rect(s,5.4,y,5.5*value,.44,col,True)
    txt(s,f'{value*100:.1f}%',11.12,y-.02,1.4,.5,25,col,True)
txt(s,'At a 191-byte gap: 9.4% local → 83.9% delta',1.0,6.15,11.3,.5,22,TEAL,True)
source(s,'Oct. 22 matched ablation: 3 seeds, 384 prompts, 1,536 updates; 32,752 vs 34,321 parameters. Narrow key/value grammar.')
note(s,55,'The clearest earlier architectural success is persistent delta memory. The model sees a key written twice and must retrieve its latest value despite distractors. With explicit training of binding, copy and writes, delta memory reaches 95.7 percent, compared with 54.9 percent for the supervised local control. The local attention window is only sixteen bytes. At the longest 191-byte gap, the difference is 83.9 versus 9.4 percent. The delta model is slightly larger and uses an additional write-gate label. This supports that combined memory recipe within a narrow grammar, not unrestricted memory or general reasoning.','Local results: reports/rewt_architecture_20261022/summary.json, memory_ablation and memory_ablation_groups\nDelta-rule background: https://arxiv.org/abs/2406.06484')

# 8
s=slide('Earlier result · Refinement bottleneck','Lower energy did not reliably fix wrong answers','Removing refinement from frozen language-model checkpoints did not reduce complete-answer accuracy.')
txt(s,'Complete-answer accuracy',.8,2.7,5.15,.45,20,INK,True)
for x,label,col in [(6.25,'Normal rounds',TEAL),(9.65,'Zero rounds',GRAY)]:txt(s,label,x,2.72,2.8,.42,18,col,True)
for y,label,a,b in [(3.55,'GRU control',34.16,34.98),(4.55,'Delta variant',35.39,35.39)]:
    txt(s,label,.9,y,4.5,.5,25,INK)
    txt(s,f'{a:.2f}%',6.25,y,2.8,.5,27,TEAL,True)
    txt(s,f'{b:.2f}%',9.65,y,2.8,.5,27,GRAY,True)
rect(s,.9,5.75,11.5,.8,INK,True)
txt(s,'The old loop optimizes the same landscape without rereading the facts.',1.12,5.98,11.06,.4,21,WHITE,True,PP_ALIGN.CENTER)
source(s,'Oct. 22 validation-only intervention: 3 seeds per variant, 162 prompts/model; includes arithmetic, rules and recall.')
note(s,50,'The core limitation was more revealing than an average score. On frozen checkpoints, removing refinement did not hurt complete-answer accuracy. For the GRU control, the zero-round policy was slightly better; for delta, it tied. Refinement did improve cross-entropy, but that is not the same as correcting whole answers. The landscape was conditioned once on the prefix. Extra optimization rounds did not perform another fact lookup. That observation motivated changing the information flow rather than simply adding more rounds. These figures are validation interventions, not new test scores.','reports/rewt_architecture_20261022/refinement_contribution.json and FINDINGS.md')

# 9
s=slide('New architecture · Opt-in structured head','Each predicted state drives another memory read','A learned byte encoder supplies keys, values, the starting query and the read count.')
for j,(a,b) in enumerate([('A','B'),('B','C'),('C','D')]):
    x=.8+j*4.2
    rect(s,x,2.8,3.6,2.4,WHITE,True)
    txt(s,f'READ {j+1}',x+.2,3.02,3.2,.3,13,TEAL,True)
    txt(s,f'query {a} → retrieve {b}',x+.2,3.65,3.2,.5,22,INK,True,PP_ALIGN.CENTER)
    txt(s,'refine this read’s state',x+.18,4.48,3.24,.4,17,GRAY,False,PP_ALIGN.CENTER)
    if j<2:arrow(s,x+3.7,3.84,.4,.25)
rect(s,.95,5.65,11.4,.85,LIGHT,True)
txt(s,'Input: raw bytes     Output: predicted nodes + predicted length',1.18,5.84,10.94,.45,22,INK,True,PP_ALIGN.CENTER)
source(s,'No symbolic parser or solver at inference. Each read defines a new energy; descent is guaranteed only within that read.')
note(s,65,'The new module uses a learned extractor over raw bytes. It predicts fact keys and values, a starting query and the requested number of reads. The current query softly attends to memory. The predicted result becomes the next query. For the example, that is A to B, then B to C, then C to D. Energy refinement can operate inside each read, but changing the retrieved context creates a new energy; there is no single descent guarantee across the whole chain. This is a structured reasoning head. It decodes intermediate nodes and their length, then serializes the trace. It is not free-form language generation. Parsers and solutions are used only to supply training labels and score results afterward.','rewt/reasoning.py; scripts/study_rewt_reasoning.py\nIllustrative flow, not a claim about this example’s measured prediction.')

# 10
s=slide('Fresh results · October 26','Repeated reads help; the wave benefit is still unproven','Complete trace + answer on 960 fresh rule prompts, across three cold starts.')
for y,label,key,col in [(2.85,'Fixed query + energy','fixed_query_energy',GRAY),(4.0,'Repeated reads + energy','reread_energy',TEAL),(5.15,'Repeated reads, zero energy rounds','reread_no_energy',INK)]:
    value=NEW['arms'][key]['test']['mean']
    txt(s,label,.82,y,5.25,.55,21,INK)
    rect(s,6.28,y+.06,4.75,.40,LIGHT,True);rect(s,6.28,y+.06,4.75*value,.40,col,True)
    txt(s,f'{value*100:.1f}%',11.25,y-.02,1.35,.55,25,col,True)
txt(s,'Both repeated-read arms also solve 480/480 prompts with eight distractor edges.',.87,6.38,11.8,.5,19,TEAL,True)
source(s,'68,502 allocated parameters; 6,000 training identities; 800 updates/fit. Same batches and labels; zero-round arm has 60,950 active parameters.')
note(s,65,f'The main fresh test contains 960 prompts, representing 480 underlying problems in familiar and two withheld phrasings. The fixed-query control averages {NEW["arms"]["fixed_query_energy"]["test"]["mean"]*100:.1f} percent. Every repeated-read seed scores 100 percent, both with and without energy refinement. Both also solve all 480 extra prompts with eight distractor edges instead of the two used in training. The important attribution is that changing the query helps here; the waves show no accuracy advantage. Removing energy from the trained energy checkpoints also preserves the perfect score. These are matched ablations of a new structured head, not a comparison against a Transformer or the old autoregressive model.','reports/rewt_reasoning_20261026/summary.json, test_results.json and test_plan.json\nAll 18,720 model/condition output rows were rechecked against the saved tasks.')

# 11
s=slide('Scope','A useful result with a clear boundary')
card(s,'DEMONSTRATED','Bounded rule composition','One to three hops.\nKnown letter vocabulary.\nShuffled facts and more distractors.',.75,2.55,5.7,3.35,TEAL)
card(s,'STILL OPEN','Broader reasoning','Natural language and arithmetic.\nLonger or branching programs.\nA causal benefit from wave refinement.',6.85,2.55,5.7,3.35,ORANGE)
txt(s,'Wording is still fragile: 43.1% on the two withheld validation phrasings.',.87,6.25,11.8,.55,20,ORANGE,True)
source(s,'The 100% test score is limited to its specified templates. Validation contains two different withheld templates and remains weaker.')
note(s,50,f'One hundred percent needs its scope stated immediately. These are acyclic letter-rule chains with at most three steps, a known vocabulary and a small set of grammars. The model is strongly biased toward differentiable lookup and receives explicit intermediate supervision. It does not establish fluent text or improved arithmetic. Wording transfer remains fragile: mean validation accuracy is {NEW["arms"]["reread_energy"]["validation_mean"]*100:.1f} percent on a split containing different withheld phrasings, even though the test phrasings happen to be easier for this architecture. The old language model is not silently replaced by this specialized head.','Study scope and exact templates: scripts/study_rewt_reasoning.py; rewt/curriculum.py\nValidation outputs and selection scores: reports/rewt_reasoning_20261026/')

# 12
s=slide('Takeaway','Make refinement earn its computational cost')
for j,(a,b) in enumerate([('Teach correction','Reward actual answer repairs'),('Integrate rereading','Bring the loop into the language model'),('Compare fairly','Measure quality at matched compute')]):
    y=2.62+j*1.05
    circle(s,str(j+1),.9,y,.68,TEAL)
    txt(s,a,1.9,y+.03,4.4,.48,25,INK,True)
    txt(s,b,6.25,y+.05,6.0,.5,21,GRAY)
rect(s,.9,6.05,11.55,.7,INK,True)
txt(s,'Memory and repeated reads help. Lower energy alone is insufficient.',1.1,6.25,11.15,.35,21,WHITE,True,PP_ALIGN.CENTER)
source(s,'Research next steps; these are not completed general-language results. Questions welcome.')
note(s,50,'The next objective is to make refinement earn its cost. That means teaching it to repair an actual wrong answer, integrating the repeated-read mechanism into a general language model, and comparing quality at matched compute with strong Transformer controls. The current contribution is an implemented, audited research system with positive and negative evidence. Persistent memory and repeated reads work on bounded tasks. Energy descent is reliable numerically, but it is not by itself evidence of reasoning. That distinction is the main lesson I would like you to take away.','Future directions grounded in reports/rewt_architecture_20261022/FINDINGS.md and reports/rewt_reasoning_20261026/FINDINGS.md')

# 13 appendix
s=slide('Appendix A · Math','The within-read descent contract','For one read, the context-dependent coefficients stay fixed.')
txt(s,'E(δ) = ½ δᵀ P δ + Σ a_k cos(ω_kᵀδ + φ_k)',.95,2.9,11.5,.75,31,INK,True,PP_ALIGN.CENTER)
txt(s,'δ = candidate − anchor       P is positive diagonal',1.0,4.0,11.3,.5,23,TEAL,False,PP_ALIGN.CENTER)
txt(s,'δ′ = δ − η ∇E(δ)       accept only if E(δ′) < E(δ)',.98,5.05,11.4,.65,27,INK,True,PP_ALIGN.CENTER)
txt(s,'A new memory read changes the anchor and amplitudes. Do not compare energies across reads.',.98,6.2,11.35,.65,20,ORANGE,True)
source(s,'Formula for the new structured head. The earlier language-model core also has quadratic compatibility terms.')
note(s,0,'Backup explanation: delta is the displacement from a new read’s anchor. Positive diagonal precision provides quadratic stabilization, and bounded Fourier modes add learned preferences. The analytic derivative is differentiated through during training. Candidate steps are checked under the same coefficients within each read. Since the next read changes those coefficients, numerical descent does not apply across reads. This formula describes the new head specifically; the older language-model energy also includes block-compatibility penalties.','rewt/reasoning.py: energy(), gradient(), refine(). Numerical comparison: tests_rewt/test_reasoning.py')

# 14
s=slide('Appendix B · Cost and ablation','The plain read loop is sufficient on this test','All arms: three seeds, identical 800-update batches and supervision.')
labels=[('Arm',.9,5.55),('Active parameters',6.0,3.0),('Training CPU / fit',9.15,3.3)]
for label,x,w in labels:txt(s,label,x,2.75,w,.45,18,TEAL,True)
for i,(label,key) in enumerate([('Fixed query + energy','fixed_query_energy'),('Repeated reads + energy','reread_energy'),('Repeated reads, zero energy rounds','reread_no_energy')]):
    y=3.55+i*.9;a=NEW['arms'][key]
    txt(s,label,.9,y,5.2,.55,21,INK)
    txt(s,f'{a["active_parameters"]:,}',6.0,y,2.8,.5,22,INK,True)
    txt(s,f'{a["training_cpu_mean"]:.1f} s',9.15,y,3.1,.5,22,INK,True)
txt(s,'Frozen intervention: disable energy → 100%; freeze queries → 33.3%.',.95,6.3,11.4,.5,21,ORANGE,True)
source(s,'Single-thread process time on a shared CPU host, not a GPU/FLOP comparison. Setup, validation and I/O excluded from fit timing.')
note(s,0,'Backup cost detail: the zero-energy arm has the same 68,502 allocated parameters, but 60,950 receive gradients. The energy-only parameters are inactive when no refinement is executed. Mean training process times include batch preparation and optimization, excluding validation, initialization and checkpoint I/O. The frozen-policy intervention uses the already trained energy models: turning off energy leaves all test answers correct, while freezing the query leaves only one-hop tasks correct. A zero-width bootstrap interval for the tied perfect scores must not be interpreted as proof of universal equality.','reports/rewt_reasoning_20261026/summary.json\nTraining CPU totals: 276.95 seconds over nine main fits; smoke runs and other work are additional.')

# 15
s=slide('Appendix C · References','Ideas are credited; project results are local')
refs=[('Vaswani et al. · Attention Is All You Need · 2017','https://arxiv.org/abs/1706.03762'),
      ('Hoover et al. · Energy Transformer · 2023','https://arxiv.org/abs/2302.07253'),
      ('Yang et al. · Parallelizing Linear Transformers with the Delta Rule · 2024','https://arxiv.org/abs/2406.06484')]
for i,(label,url) in enumerate(refs):
    sh=txt(s,label,.85,2.62+i*.72,11.65,.55,21,INK)
    for run in sh.text_frame.paragraphs[0].runs:run.hyperlink.address=url
txt(s,'Project evidence',.85,5.05,11,.45,20,TEAL,True)
txt(s,'Architecture, memory and numerical audits · October 22\nLearned rereading, controls and raw outputs · October 26',.85,5.68,11.7,.95,21,INK)
source(s,'Local source reports: reports/rewt_architecture_20261022/FINDINGS.md and reports/rewt_reasoning_20261026/FINDINGS.md')
note(s,0,'Background references are linked on this slide. The project’s numerical and quality results are from the local recorded experiments, not borrowed from those papers. The new repeated-read head is an explicitly constrained research extension and should not be presented as a reproduction of the published Energy Transformer or DeltaNet systems.','https://arxiv.org/abs/1706.03762\nhttps://arxiv.org/abs/2302.07253\nhttps://arxiv.org/abs/2406.06484\nreports/rewt_architecture_20261022/FINDINGS.md\nreports/rewt_reasoning_20261026/FINDINGS.md')

assert len(prs.slides)==15
assert sum(n['seconds'] for n in notes)==600
prs.save(HERE/'REWT_class_presentation.pptx')
(HERE/'speaker_notes.txt').write_text('REWT · TEN-MINUTE REHEARSAL SCRIPT\n12 main slides, 3 backup slides. Timings total 10:00.\n\n'+'\n\n' .join(n['text'] for n in notes)+'\n')
(HERE/'slide_manifest.json').write_text(json.dumps(dict(main_slides=12,appendix_slides=3,planned_seconds=600,
    slides=[dict(slide=n['slide'],seconds=n['seconds']) for n in notes],
    editable=True,source_results=['reports/rewt_architecture_20261022/summary.json','reports/rewt_reasoning_20261026/summary.json']),indent=2)+'\n')
print('Built 15 editable slides with speaker notes; 12-slide main talk = 10:00.')
